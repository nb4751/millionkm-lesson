"""End-to-end request tests for authorization, uploads and reports."""
import hashlib
import io
import json
import zipfile
from datetime import timedelta

import pytest
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import Client
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from training.models import Activity, Athlete, AthleteAccessLink, DataRequest, FTPRecord, Invitation

pytestmark = pytest.mark.django_db


@pytest.fixture(autouse=True)
def explicit_auth_mode(settings):
    # Request-level permission tests exercise the secure remote/account mode.
    settings.APP_LOCAL_COACH = False
    settings.APP_OPEN_PORTAL = False


@pytest.fixture
def roster(settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path / "uploads"
    settings.APP_DATA_DIR = tmp_path
    settings.DATA_DIR = tmp_path
    User = get_user_model()
    coach = User.objects.create_user("coach", password="Coach-test-pass-932!", is_staff=True)
    user_a = User.objects.create_user("athlete-a", password="Rider-test-pass-932!")
    user_b = User.objects.create_user("athlete-b", password="Rider-test-pass-933!")
    a = Athlete.objects.create(name="격리 가상 A", code="TEST-A", user=user_a, weight_kg=60)
    b = Athlete.objects.create(name="격리 가상 B", code="TEST-B", user=user_b, weight_kg=70)
    FTPRecord.objects.create(athlete=a, value=250, kind="coach", date=timezone.localdate())
    return coach, a, b


def activity(athlete, **kwargs):
    defaults = {
        "title": f"{athlete.name} 비공개 라이딩", "start_time": timezone.now(),
        "duration_seconds": 120, "distance_km": 1.5,
        "source_hash": hashlib.sha256(athlete.code.encode()).hexdigest(),
        "metrics": {"avg_power": 200, "max_power": 240, "tss": 10,
                    "power_curve": {"5": 230, "60": 220, "300": None, "1200": None}},
    }
    defaults.update(kwargs)
    return Activity.objects.create(athlete=athlete, **defaults)


def csv_file(name="ride.csv"):
    text = "timestamp,power,heart_rate,cadence\n"
    for second in range(61):
        text += f"2026-10-01T09:{second // 60:02d}:{second % 60:02d}+09:00,200,140,90\n"
    return SimpleUploadedFile(name, text.encode(), content_type="text/csv")


def stream_bytes(response):
    if response.streaming:
        return b"".join(response.streaming_content)
    return response.content


def test_anonymous_cannot_read_dashboard_or_report(client):
    for route in ["dashboard", "export", "upload", "compare"]:
        response = client.get(reverse(route))
        assert response.status_code == 302
        assert reverse("login") in response.url


def test_athlete_cannot_read_another_athlete_or_activity(client, roster):
    coach, a, b = roster
    private = activity(b)
    client.force_login(a.user)
    assert client.get(reverse("athlete_detail", args=[b.pk])).status_code == 404
    assert client.get(reverse("activity_detail", args=[private.pk])).status_code == 404
    own = activity(a)
    assert client.get(reverse("activity_detail", args=[own.pk])).status_code == 200


def test_dashboard_and_export_exclude_other_athletes(client, roster):
    coach, a, b = roster
    activity(a)
    private = activity(b)
    client.force_login(a.user)
    response = client.get(reverse("dashboard"))
    assert response.status_code == 200
    assert b.name.encode() not in response.content
    # An adversarial selection must never override object-level access control.
    response = client.get(reverse("export"), {"athletes": [b.pk], "athlete": b.pk})
    assert response.status_code in [200, 403, 404]
    if response.status_code == 200:
        workbook = load_workbook(io.BytesIO(stream_bytes(response)))
        all_values = [str(cell.value) for sheet in workbook for row in sheet for cell in row]
        assert b.name not in all_values
        assert private.title not in all_values


def test_coach_reads_all_athletes_and_report_sheets(client, roster):
    coach, a, b = roster
    activity(a)
    activity(b)
    client.force_login(coach)
    assert client.get(reverse("athlete_detail", args=[b.pk])).status_code == 200
    response = client.get(reverse("export"))
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(stream_bytes(response)))
    assert workbook.sheetnames == ["선수목록", "활동요약", "주간훈련량", "파워커브", "FTP이력", "파워존", "인터벌분석", "비교리포트"]
    activity_values = [cell.value for row in workbook["활동요약"] for cell in row]
    assert a.name in activity_values and b.name in activity_values
    assert 200 in activity_values
    assert 10 in activity_values


def test_upload_is_scoped_and_duplicates_rejected(client, roster):
    coach, a, b = roster
    client.force_login(a.user)
    response = client.post(reverse("upload"), {"athlete": b.pk, "file": csv_file(), "kind": "personal"})
    assert response.status_code in [200, 302, 403, 404]
    assert not Activity.objects.filter(athlete=b).exists()
    # An authenticated rider may upload only into their own account.
    response = client.post(reverse("upload"), {"file": csv_file(), "athlete": a.pk, "kind": "personal"})
    assert response.status_code in [200, 302]
    assert Activity.objects.filter(athlete=a).count() == 1
    response = client.post(reverse("upload"), {"file": csv_file(), "athlete": a.pk, "kind": "personal"})
    assert response.status_code in [200, 302]
    assert Activity.objects.filter(athlete=a).count() == 1


@pytest.mark.parametrize("filename,content", [("payload.exe", b"MZmalware"), ("fake.fit", b"this is not a FIT activity"), ("fake.csv", b"PK\x03\x04zip-not-csv")])
def test_upload_rejects_untrusted_content(client, roster, filename, content):
    _, a, _ = roster
    client.force_login(a.user)
    response = client.post(reverse("upload"), {"file": SimpleUploadedFile(filename, content), "athlete": a.pk, "kind": "personal"})
    assert response.status_code in [200, 400]
    assert Activity.objects.count() == 0


def test_server_enforces_upload_size_limit(client, roster, settings):
    _, a, _ = roster
    settings.MAX_UPLOAD_BYTES = 32
    client.force_login(a.user)
    response = client.post(reverse("upload"), {"file": csv_file(), "athlete": a.pk, "kind": "personal"})
    assert response.status_code in [200, 400, 413]
    assert Activity.objects.count() == 0


def token_for(athlete, token, expires=None, used=None):
    return Invitation.objects.create(
        athlete=athlete, token_hash=hashlib.sha256(token.encode()).hexdigest(),
        expires_at=expires or timezone.now() + timedelta(hours=1), used_at=used,
    )


@pytest.mark.parametrize("state", ["expired", "used"])
def test_invalid_invitation_cannot_activate_account(client, roster, state):
    _, _, b = roster
    token = "invitation-test-secret-token"
    token_for(b, token, expires=timezone.now() - timedelta(seconds=1) if state == "expired" else None,
              used=timezone.now() if state == "used" else None)
    response = client.get(reverse("activate", args=[token]))
    assert response.status_code in [400, 404, 410]
    response = client.post(reverse("activate", args=[token]), {"username": "intruder", "password1": "Safe-password-329!", "password2": "Safe-password-329!", "consent": True})
    assert response.status_code in [400, 404, 410]
    assert not get_user_model().objects.filter(username="intruder").exists()


def test_real_invitation_is_single_use(client, roster):
    coach, _, _ = roster
    fresh = Athlete.objects.create(name="초대 가상선수", code="TEST-INVITE")
    token = "single-use-invitation-test-token"
    invitation = token_for(fresh, token)
    response = client.post(reverse("activate", args=[token]), {
        "username": "new-rider", "password1": "New-rider-password-782!", "password2": "New-rider-password-782!",
        "consent": True, "consent_note": "가상 사용자 시험 동의",
    })
    assert response.status_code == 302
    invitation.refresh_from_db()
    fresh.refresh_from_db()
    assert invitation.used_at is not None
    assert fresh.user is not None
    assert fresh.user.check_password("New-rider-password-782!")
    assert client.get(reverse("activate", args=[token])).status_code in [400, 404, 410]


def test_csrf_is_required_for_upload_and_login(roster):
    _, a, _ = roster
    client = Client(enforce_csrf_checks=True)
    assert client.post(reverse("login"), {"username": "athlete-a", "password": "Rider-test-pass-932!"}).status_code == 403
    client.force_login(a.user)
    assert client.post(reverse("upload"), {"file": csv_file()}).status_code == 403


def test_local_coach_enters_without_login_but_external_client_does_not(settings):
    settings.APP_LOCAL_COACH = True
    local = Client(REMOTE_ADDR="127.0.0.1")
    response = local.get(reverse("dashboard"))
    assert response.status_code == 200
    owner = get_user_model().objects.get(username="__local_coach__")
    assert owner.is_staff
    assert not owner.has_usable_password()
    assert local.get(reverse("athlete_add")).status_code == 200
    external = Client(REMOTE_ADDR="203.0.113.10")
    response = external.get(reverse("dashboard"))
    assert response.status_code == 302
    assert reverse("login") in response.url


def access_token(athlete, token, **kwargs):
    return AthleteAccessLink.objects.create(
        athlete=athlete, token_hash=hashlib.sha256(token.encode()).hexdigest(),
        expires_at=kwargs.pop("expires_at", timezone.now() + timedelta(days=1)), **kwargs,
    )


def test_personal_link_needs_no_password_and_stays_in_athlete_scope(client, roster, settings):
    _, a, b = roster
    settings.APP_LOCAL_COACH = True
    own = activity(a)
    other = activity(b)
    token = "only-this-athlete-secret-token"
    link = access_token(a, token)
    response = client.get(reverse("shared_workspace", args=[token]), follow=True)
    assert response.status_code == 200
    assert client.session["athlete_link_id"] == link.pk
    assert client.get(reverse("activity_detail", args=[own.pk])).status_code == 200
    assert client.get(reverse("activity_detail", args=[other.pk])).status_code == 404
    assert client.get(reverse("athlete_add")).status_code == 404
    assert client.post(reverse("backup")).status_code == 404


def test_personal_link_revocation_ends_existing_session(client, roster):
    _, a, _ = roster
    token = "revocation-test-secret-token"
    link = access_token(a, token)
    assert client.get(reverse("shared_workspace", args=[token])).status_code == 302
    link.revoked_at = timezone.now()
    link.save(update_fields=["revoked_at"])
    assert client.get(reverse("athlete_detail", args=[a.pk])).status_code == 403
    assert client.get(reverse("shared_workspace", args=[token])).status_code == 404


def test_replacement_personal_link_recovers_revoked_session(client, roster):
    _, a, _ = roster
    old = access_token(a, "old-personal-link")
    client.get(reverse("shared_workspace", args=["old-personal-link"]))
    old.revoked_at = timezone.now()
    old.save(update_fields=["revoked_at"])
    new = access_token(a, "replacement-personal-link")
    response = client.get(reverse("shared_workspace", args=["replacement-personal-link"]), follow=True)
    assert response.status_code == 200
    assert client.session["athlete_link_id"] == new.pk


@pytest.mark.parametrize("state", ["expired", "revoked", "disabled"])
def test_invalid_personal_links_are_rejected(client, roster, state):
    _, a, _ = roster
    options = {}
    if state == "expired":
        options["expires_at"] = timezone.now() - timedelta(seconds=1)
    if state == "revoked":
        options["revoked_at"] = timezone.now()
    if state == "disabled":
        a.is_active = False
        a.save(update_fields=["is_active"])
    access_token(a, "invalid-personal-link", **options)
    assert client.get(reverse("shared_workspace", args=["invalid-personal-link"])).status_code == 404


def test_shared_portal_allows_names_reading_and_upload_without_login(roster, settings):
    _, a, b = roster
    settings.APP_OPEN_PORTAL = True
    external = Client(REMOTE_ADDR="203.0.113.20")
    own = activity(a)
    other = activity(b)
    response = external.get(reverse("dashboard"))
    assert response.status_code == 200
    assert a.name.encode() in response.content
    assert b.name.encode() in response.content
    assert external.get(reverse("activity_detail", args=[own.pk])).status_code == 200
    assert external.get(reverse("activity_detail", args=[other.pk])).status_code == 200
    assert external.get(reverse("compare")).status_code == 200
    assert external.get(reverse("export")).status_code == 200
    response = external.post(reverse("upload"), {"athlete": b.pk, "file": csv_file(), "kind": "personal"})
    assert response.status_code == 302
    assert Activity.objects.filter(athlete=b).count() == 2
    assert external.post(reverse("backup")).status_code == 404
    assert external.get(reverse("athlete_add")).status_code == 404
    # A public visitor cannot erase records via forged POSTs.
    count = Activity.objects.filter(athlete=b).count()
    response = external.post(reverse("activity_delete", args=[other.pk]))
    assert response.status_code in [302, 403, 404]
    assert Activity.objects.filter(athlete=b).count() == count


def test_shared_portal_compare_does_not_expose_inactive_names(roster, settings):
    _, active, inactive = roster
    settings.APP_OPEN_PORTAL = True
    inactive.is_active = False
    inactive.save(update_fields=["is_active"])
    external = Client(REMOTE_ADDR="203.0.113.21")
    response = external.get(reverse("compare"))
    assert response.status_code == 200
    assert active.name.encode() in response.content
    assert inactive.name.encode() not in response.content
    assert inactive.pk not in [a.pk for a in response.context["athletes"]]


def test_bulk_summary_storage_limit_counts_each_original_copy(roster, settings, monkeypatch):
    _, a, _ = roster
    settings.APP_OPEN_PORTAL = True
    monkeypatch.setenv("APP_MAX_STORAGE_MB", "1")
    start = timezone.now()
    parsed = {"activities": [
        {"title": "요약 A", "start_time": start, "duration_seconds": 600, "samples": []},
        {"title": "요약 B", "start_time": start + timedelta(days=1), "duration_seconds": 600, "samples": []},
    ], "warnings": []}
    monkeypatch.setattr("training.views.parse_file", lambda *args, **kwargs: parsed)
    # A single file fits the 1MiB quota; storing it once for each activity does not.
    content = b"summary" * 90000
    external = Client(REMOTE_ADDR="203.0.113.22")
    response = external.post(reverse("upload"), {
        "athlete": a.pk, "kind": "personal",
        "file": SimpleUploadedFile("many-activities.csv", content, content_type="text/csv"),
    })
    assert response.status_code == 200
    assert response.context["form"].errors
    assert Activity.objects.count() == 0
    assert not any(p.is_file() for p in settings.MEDIA_ROOT.rglob("*"))


def test_shared_portal_cannot_download_original_with_possible_gps(roster, settings):
    from django.core.files.base import ContentFile
    coach, a, _ = roster
    settings.APP_OPEN_PORTAL = True
    ride = activity(a)
    ride.original_file.save("private.fit", ContentFile(b"sensitive-original-gps-data"))
    external = Client(REMOTE_ADDR="203.0.113.23")
    assert external.get(reverse("activity_detail", args=[ride.pk])).status_code == 200
    assert external.get(reverse("original_file", args=[ride.pk])).status_code == 404
    private = Client()
    private.force_login(coach)
    response = private.get(reverse("original_file", args=[ride.pk]))
    assert response.status_code == 200
    assert stream_bytes(response) == b"sensitive-original-gps-data"


def test_shared_portal_request_uses_selected_athlete(roster, settings):
    _, a, b = roster
    settings.APP_OPEN_PORTAL = True
    external = Client(REMOTE_ADDR="203.0.113.24")
    response = external.get(reverse("account_request"))
    assert response.status_code == 200
    choices = response.context["form"].fields["athlete"].queryset
    assert set(choices.values_list("pk", flat=True)) == {a.pk, b.pk}
    response = external.post(reverse("account_request"), {
        "athlete": b.pk, "kind": "delete", "note": "선택한 선수의 확인용 삭제 요청",
    })
    assert response.status_code == 302
    request = DataRequest.objects.get()
    assert request.athlete_id == b.pk
    assert request.status == "pending"
    assert Athlete.objects.filter(pk=b.pk).exists()


def test_coach_functions_forbidden_for_athlete(client, roster):
    _, a, b = roster
    client.force_login(a.user)
    for name, args in [("athlete_add", []), ("athlete_invite", [b.pk]), ("backup", [])]:
        response = client.post(reverse(name, args=args))
        assert response.status_code in [403, 404]


def test_seed_demo_is_repeatable_and_does_not_create_accounts(settings, tmp_path):
    settings.APP_DATA_DIR = tmp_path
    settings.DATA_DIR = tmp_path
    call_command("seed_demo", stdout=io.StringIO())
    assert Athlete.objects.filter(is_demo=True).count() == 3
    assert Activity.objects.count() == 18
    assert get_user_model().objects.count() == 0
    call_command("seed_demo", stdout=io.StringIO())
    assert Athlete.objects.count() == 3
    assert Activity.objects.count() == 18
    assert (tmp_path / "samples" / "가상_라이딩_업로드_예제.csv").is_file()


def test_backup_archive_has_checksums_and_no_environment_secrets(settings, tmp_path, monkeypatch):
    import sqlite3
    db = tmp_path / "source.sqlite3"
    with sqlite3.connect(db) as connection:
        connection.execute("CREATE TABLE example (id INTEGER)")
        connection.execute("INSERT INTO example VALUES (42)")
    # The command reads configuration directly; keep the test runner's own live
    # database connection untouched while supplying a real backup source file.
    monkeypatch.setitem(settings.DATABASES, "default", {"ENGINE": "django.db.backends.sqlite3", "NAME": str(db)})
    settings.MEDIA_ROOT = tmp_path / "uploads"
    settings.MEDIA_ROOT.mkdir()
    (settings.MEDIA_ROOT / "original.csv").write_text("power\n200\n")
    (tmp_path / ".env").write_text("APP_SECRET_KEY=not-for-backup")
    destination = tmp_path / "backup.zip"
    call_command("backup_data", output=str(destination), stdout=io.StringIO())
    with zipfile.ZipFile(destination) as archive:
        assert set(archive.namelist()) == {"db.sqlite3", "uploads/original.csv", "manifest.json"}
        manifest = json.loads(archive.read("manifest.json"))
        for filename, expected in manifest["sha256"].items():
            assert hashlib.sha256(archive.read(filename)).hexdigest() == expected
        restored = tmp_path / "restored.sqlite3"
        restored.write_bytes(archive.read("db.sqlite3"))
    with sqlite3.connect(restored) as connection:
        assert connection.execute("SELECT id FROM example").fetchone() == (42,)
