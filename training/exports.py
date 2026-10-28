from collections import defaultdict
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, Reference
from django.utils import timezone
from .analysis import ftp_zones

def safe_cell(value):
    # Guard spreadsheet formula injection in all user-originated cells.
    if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@')):
        return "'" + value
    return value

def workbook_bytes(athletes, activities, comparisons):
    wb = Workbook()
    wb.remove(wb.active)
    names = ['선수목록','활동요약','주간훈련량','파워커브','FTP이력','파워존','인터벌분석','비교리포트']
    sheets = {n:wb.create_sheet(n) for n in names}
    def row(sheet, values):
        sheets[sheet].append([safe_cell(v) for v in values])
    row('선수목록',['선수 ID','이름','그룹','체중(kg)','체중 측정일','가상 데이터'])
    row('활동요약',['선수 ID','선수','활동','시작(KST)','시간(초)','거리(km)','평균파워(W)','NP(W)','IF','TSS','FTP 기준(W)','파워 측정비율'])
    row('주간훈련량',['선수 ID','주 시작(KST)','시간(h)','거리(km)','TSS','부하 미측정 활동수'])
    row('파워커브',['선수 ID','활동','시작(KST)','구간(초)','최대평균파워(W)'])
    row('FTP이력',['선수 ID','적용일','FTP(W)','출처','방법','기준 사용'])
    row('파워존',['선수 ID','기준 FTP(W)','영역','최저(W)','최고(W)'])
    row('인터벌분석',['선수 ID','활동','시작(초)','종료(초)','목표(W)','평균(W)','달성률(%)'])
    row('비교리포트',['선수 ID','선수','FTP(W)','FTP(W/kg)','5초(W)','1분(W)','5분(W)','20분(W)','시간(h)','TSS','활동수'])
    weekly = defaultdict(lambda:[0,0,0,0,0])
    for a in athletes:
        row('선수목록',[a.code,a.name,a.group,a.weight_kg,a.weight_date.isoformat() if a.weight_date else None,a.is_demo])
        for f in a.ftp_records.order_by('date','id'):
            row('FTP이력',[a.code,f.date.isoformat(),f.value,f.get_kind_display(),f.method,f.is_active])
        for z in ftp_zones(a.ftp) if a.ftp else []:
            row('파워존',[a.code,a.ftp,z['name'],z['min'],z['max']])
    for a in activities:
        m=a.metrics
        date=timezone.localtime(a.start_time)
        stamp=date.strftime('%Y-%m-%d %H:%M:%S')
        row('활동요약',[a.athlete.code,a.athlete.name,a.title,stamp,a.duration_seconds,a.distance_km,m.get('avg_power'),m.get('np'),m.get('if'),m.get('tss'),m.get('ftp_used'),m.get('power_coverage')])
        week=(date.date()-__import__('datetime').timedelta(days=date.weekday())).isoformat()
        w=weekly[(a.athlete.code,week)]; w[0]+=a.duration_seconds/3600; w[1]+=a.distance_km or 0
        if m.get('tss') is None: w[3]+=1
        else: w[2]+=m['tss']; w[4]+=1
        for seconds,value in m.get('power_curve',{}).items():
            row('파워커브',[a.athlete.code,a.title,stamp,int(seconds),value])
        for i in a.intervals.all():
            row('인터벌분석',[a.athlete.code,a.title,i.start_seconds,i.end_seconds,i.target_power,i.metrics.get('avg_power'),i.metrics.get('achievement')])
    for (code,week),w in sorted(weekly.items()):
        row('주간훈련량',[code,week,w[0],w[1],w[2] if w[4] else None,w[3]])
    for r in comparisons:
        row('비교리포트',[r['athlete'].code,r['athlete'].name,r['ftp'],r['wkg'],r['mmp5'],r['mmp60'],r['mmp300'],r['mmp1200'],r['hours'],r['tss'],r['activity_count']])
    for ws in wb:
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:
            cell.font=Font(color='FFFFFF',bold=True); cell.fill=PatternFill('solid',fgColor='172D34')
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width=min(34,max(14,max(len(str(c.value or '')) for c in col)+3))
        for r in ws.iter_rows(min_row=2):
            for c in r:
                if isinstance(c.value,float): c.number_format='0.00'
    ws=sheets['비교리포트']
    if ws.max_row>1:
        chart=BarChart(); chart.title='선수별 FTP (W)'
        chart.add_data(Reference(ws,min_col=3,min_row=1,max_row=ws.max_row),titles_from_data=True)
        chart.set_categories(Reference(ws,min_col=2,min_row=2,max_row=ws.max_row)); ws.add_chart(chart,'N2')
    out=BytesIO(); wb.save(out); return out.getvalue()
